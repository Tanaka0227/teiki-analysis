import csv
import os
import glob
import webbrowser
import subprocess
from collections import defaultdict

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

def find_csv():
    csvs = glob.glob(os.path.join(SCRIPT_DIR, "*.csv"))
    if not csvs:
        raise FileNotFoundError("CSVファイルが見つかりません。フォルダにCSVを入れてください。")
    csvs.sort(key=os.path.getmtime, reverse=True)
    return csvs[0]

def find_master():
    path = os.path.join(SCRIPT_DIR, "原価マスタ.xlsx")
    if not os.path.exists(path):
        raise FileNotFoundError("原価マスタ.xlsx が見つかりません。")
    return path

def load_csv(path):
    for enc in ['cp932', 'utf-8', 'shift_jis']:
        try:
            with open(path, 'r', encoding=enc, errors='replace') as f:
                reader = csv.DictReader(f)
                return list(reader)
        except Exception:
            continue
    raise ValueError("CSVの読み込みに失敗しました")

def load_master(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {}
        for i, h in enumerate(headers):
            if h and i < len(row):
                d[h] = row[i]
        rows.append(d)
    return rows

def calculate(csv_rows, master_rows):
    product_map = {}
    for r in master_rows:
        code = str(r.get('商品コード', '')).strip()
        if not code:
            continue
        try:
            bags = int(r.get('個数', 1))
        except (ValueError, TypeError):
            bags = 1
        is_upsell = str(r.get('アップセル', '')).strip() == 'おまとめ'
        product_map[code] = {'bags': bags, 'is_upsell': is_upsell, 'name': str(r.get('商品名', ''))}

    teiki_to_media = {}
    teiki_n_data = defaultdict(lambda: defaultdict(lambda: {'codes': [], 'subtotal': 0}))

    for row in csv_rows:
        n_str = row.get('定期内受注（N番目）', '').strip()
        teiki = row.get('定期受注番号', '').strip()
        media = row.get('広告URLグループ名', '').strip()
        code = row.get('購入商品（商品コード）', '').strip()
        try:
            subtotal = float(row.get('小計', '0') or '0')
        except ValueError:
            subtotal = 0
        try:
            n = int(n_str)
        except (ValueError, TypeError):
            continue
        if not teiki:
            continue
        if n == 1 and media:
            teiki_to_media[teiki] = media
        teiki_n_data[teiki][n]['codes'].append(code)
        teiki_n_data[teiki][n]['subtotal'] += subtotal

    media_stats = {}
    for teiki, n_map in teiki_n_data.items():
        media = teiki_to_media.get(teiki)
        if not media:
            continue
        if media not in media_stats:
            media_stats[media] = {
                'q1_bags': 0, 'q2': 0, 'upsell_count': 0,
                'n_counts': defaultdict(set), 'n_revenue': defaultdict(float),
                'product_name': ''
            }
        ms = media_stats[media]
        for n, data in n_map.items():
            ms['n_counts'][n].add(teiki)
            ms['n_revenue'][n] += data['subtotal']
            if n == 1:
                total_bags = 0
                has_upsell = False
                for c in data['codes']:
                    prod = product_map.get(c)
                    if prod:
                        total_bags += prod['bags']
                        if prod['is_upsell']:
                            has_upsell = True
                        if not ms['product_name']:
                            ms['product_name'] = prod['name']
                    else:
                        total_bags += 1
                ms['q1_bags'] += total_bags
                ms['q2'] += 1
                if has_upsell:
                    ms['upsell_count'] += 1

    results = []
    for media, ms in media_stats.items():
        n1 = ms['q2']
        if n1 == 0:
            continue
        upsell_rate = ms['upsell_count'] / n1
        rates = []
        for i in range(1, 13):
            cnt = len(ms['n_counts'].get(i, set()))
            rates.append(cnt / n1 if cnt > 0 else None)
        rev6 = sum(ms['n_revenue'].get(n, 0) for n in range(1, 7))
        rev12 = sum(ms['n_revenue'].get(n, 0) for n in range(1, 13))
        ltv6 = round(rev6 / n1)
        ltv12 = round(rev12 / n1)
        results.append({
            'media': media,
            'product': ms['product_name'],
            'q1': ms['q1_bags'],
            'q2': n1,
            'upsell_q': ms['upsell_count'],
            'upsell_rate': upsell_rate,
            'rates': rates,
            'ltv6': ltv6,
            'ltv12': ltv12
        })

    results.sort(key=lambda x: x['q2'], reverse=True)
    return results

def pct(v):
    if v is None:
        return ''
    return f'{v * 100:.1f}%'

def build_tsv(results):
    headers = ['広告URLグループ', '商品名', 'Q1（袋数）', 'Q2（件数）', 'アップセルQ', 'アップセル率',
               'F1', 'F2', 'F3', 'F4', 'F5', 'F6', 'F7', 'F8', 'F9', 'F10', 'F11', 'F12',
               '半年LTV', '一年LTV']
    lines = ['\t'.join(headers)]
    for r in results:
        row = [r['media'], r['product'], str(r['q1']), str(r['q2']),
               str(r['upsell_q']), pct(r['upsell_rate'])]
        row += [pct(v) for v in r['rates']]
        row += [str(r['ltv6']), str(r['ltv12'])]
        lines.append('\t'.join(row))
    return '\n'.join(lines)

def build_html(results, csv_name):
    headers = ['広告URLグループ', '商品名', 'Q1（袋数）', 'Q2（件数）', 'アップセルQ', 'アップセル率',
               'F1', 'F2', 'F3', 'F4', 'F5', 'F6', 'F7', 'F8', 'F9', 'F10', 'F11', 'F12',
               '半年LTV', '一年LTV']

    total_q2 = sum(r['q2'] for r in results)

    rows_html = ''
    for r in results:
        cells = f'<td style="text-align:left">{r["media"]}</td><td style="text-align:left">{r["product"]}</td>'
        cells += f'<td>{r["q1"]}</td><td>{r["q2"]}</td><td>{r["upsell_q"]}</td><td>{pct(r["upsell_rate"])}</td>'
        for v in r['rates']:
            cells += f'<td>{pct(v)}</td>'
        cells += f'<td>{r["ltv6"]}</td><td>{r["ltv12"]}</td>'
        rows_html += f'<tr>{cells}</tr>\n'

    th_html = ''.join(f'<th>{h}</th>' for h in headers)

    return f'''<!DOCTYPE html>
<html lang="ja">
<head>
<meta charset="UTF-8">
<title>STRAC結果</title>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: 'Segoe UI', 'Meiryo', sans-serif; background: #f5f7fa; color: #333; padding: 20px; }}
  h1 {{ font-size: 22px; margin-bottom: 8px; color: #1a1a2e; }}
  .info {{ font-size: 13px; color: #888; margin-bottom: 16px; }}
  .summary {{ background: #fff; padding: 14px 18px; border-radius: 8px; margin-bottom: 16px; box-shadow: 0 1px 4px rgba(0,0,0,0.1); font-size: 14px; display: flex; gap: 24px; align-items: center; }}
  .summary span {{ font-weight: bold; color: #4a6fa5; }}
  .btn {{ background: #27ae60; color: #fff; border: none; padding: 8px 20px; border-radius: 6px; font-size: 14px; cursor: pointer; }}
  .btn:hover {{ background: #1e8c4c; }}
  .msg {{ color: #27ae60; font-size: 13px; margin-left: 10px; opacity: 0; transition: opacity 0.3s; }}
  .msg.show {{ opacity: 1; }}
  .table-wrap {{ overflow-x: auto; }}
  table {{ border-collapse: collapse; width: 100%; background: #fff; border-radius: 8px; overflow: hidden; box-shadow: 0 1px 4px rgba(0,0,0,0.1); font-size: 13px; }}
  th {{ background: #1a1a2e; color: #fff; padding: 8px 10px; text-align: center; white-space: nowrap; position: sticky; top: 0; }}
  td {{ padding: 6px 10px; text-align: center; border-bottom: 1px solid #eee; white-space: nowrap; }}
  tr:hover td {{ background: #f0f4ff; }}
</style>
</head>
<body>
<h1>媒体別STRAC算出結果</h1>
<div class="info">ソース: {csv_name}</div>
<div class="summary">
  <div>媒体数: <span>{len(results)}</span></div>
  <div>総Q2: <span>{total_q2}</span></div>
  <button class="btn" onclick="copyTSV()">クリップボードにコピー</button>
  <span class="msg" id="msg">コピーしました！</span>
</div>
<div class="table-wrap">
<table>
<thead><tr>{th_html}</tr></thead>
<tbody>
{rows_html}
</tbody>
</table>
</div>
<textarea id="tsv" style="position:absolute;left:-9999px">{build_tsv(results)}</textarea>
<script>
function copyTSV() {{
  const ta = document.getElementById('tsv');
  ta.style.position = 'static';
  ta.select();
  document.execCommand('copy');
  ta.style.position = 'absolute';
  const m = document.getElementById('msg');
  m.classList.add('show');
  setTimeout(() => m.classList.remove('show'), 2000);
}}
</script>
</body>
</html>'''

def copy_to_clipboard(text):
    try:
        subprocess.run(['clip'], input=text.encode('utf-16-le'), check=True)
        return True
    except Exception:
        return False

def main():
    print("STRAC算出ツール")
    print("=" * 40)

    csv_path = find_csv()
    csv_name = os.path.basename(csv_path)
    print(f"CSV: {csv_name}")

    master_path = find_master()
    print(f"原価マスタ: {os.path.basename(master_path)}")

    print("\n読み込み中...")
    csv_rows = load_csv(csv_path)
    print(f"  CSV: {len(csv_rows)} 行")

    master_rows = load_master(master_path)
    print(f"  原価マスタ: {len(master_rows)} 商品")

    print("\n計算中...")
    results = calculate(csv_rows, master_rows)
    print(f"  {len(results)} 媒体を検出")

    tsv = build_tsv(results)
    if copy_to_clipboard(tsv):
        print("\nTSVをクリップボードにコピーしました（スプシに貼り付け可能）")

    html = build_html(results, csv_name)
    out_path = os.path.join(SCRIPT_DIR, "strac_result.html")
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"結果HTML: {out_path}")

    webbrowser.open(f'file:///{out_path.replace(os.sep, "/")}')
    print("\nブラウザで結果を表示しました。")

    print("\n--- 上位5媒体 ---")
    for r in results[:5]:
        print(f"  {r['media']}: Q1={r['q1']}, Q2={r['q2']}, LTV6={r['ltv6']}, LTV12={r['ltv12']}")

if __name__ == '__main__':
    main()
